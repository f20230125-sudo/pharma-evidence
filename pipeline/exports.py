"""Create reconciled Excel and Power BI inputs, without invented operational KPIs."""
import csv
import json
import pathlib
import shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.chart import BarChart, Reference
from openpyxl.workbook.properties import CalcProperties

ROOT=pathlib.Path(__file__).resolve().parents[1]
out=ROOT/'public/downloads';out.mkdir(parents=True,exist_ok=True)
manifest=json.loads((ROOT/'data/manifest.json').read_text())
benchmark=json.loads((ROOT/'data/evaluation/benchmark.json').read_text(encoding='utf8'))
wb=Workbook();readme=wb.active;readme.title='Read me'
for row in [
 ['PharmaEvidence | Source and review tracker'],
 ['Snapshot','2026-10-16'],
 ['Scope','Public US sources; two MSD products; fictional drafts; human review required.'],
 ['Authority','DailyMed is manufacturer-submitted in-use labeling. Full FDA-approved-text equivalence is not asserted.'],
 ['Formulas','Dashboard formulas count the actual registry and section-change tables. No sales or time-savings claims.'],
 ['Privacy','Browser review events are not included. Export them from the app to analyze your own local decisions.'],
 ['Benchmark','Authored fixtures, not expert-adjudicated regulatory/medical ground truth.'],
 ['Source pipeline','python pipeline/build.py && python pipeline/benchmark.py && python pipeline/exports.py'],
 ['Power BI','Import source_versions.csv, section_changes.csv and exported review_events.csv. See powerbi/README.md.']
]:readme.append(row)
for title,filename in [('Versions','source_versions.csv'),('Section changes','section_changes.csv')]:
 ws=wb.create_sheet(title)
 with (ROOT/'data/marts'/filename).open(encoding='utf8',newline='') as f:
  for row in csv.reader(f):ws.append(row)
 table=Table(displayName='SourceVersions' if title=='Versions' else 'SectionChanges',ref=ws.dimensions)
 table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium4',showRowStripes=True)
 ws.add_table(table);shutil.copy2(ROOT/'data/marts'/filename,out/filename)
fixture=wb.create_sheet('Benchmark fixtures')
fixture.append(['case_id','family','split','product','claim','expected_supported','evidence_id'])
for c in benchmark['cases']:fixture.append([c['id'],c['family'],c['split'],c['product'],c['claim'],c['supported'],c['evidenceId']])
fixture.add_table(Table(displayName='ClaimFixtures',ref=fixture.dimensions))
dash=wb.create_sheet('Dashboard',1)
for r in [['SOURCE SNAPSHOT','Count'],['Genuine source versions','=COUNTA(SourceVersions[id])'],
          ['Changed sections','=COUNTIF(SectionChanges[kind],"modified")'],['Added sections','=COUNTIF(SectionChanges[kind],"added")'],
          ['Removed sections','=COUNTIF(SectionChanges[kind],"removed")'],['Unchanged sections','=COUNTIF(SectionChanges[kind],"unchanged")'],
          ['Reference-only changes','=COUNTIF(SectionChanges[kind],"reference-only")'],
          ['Authored benchmark cases','=COUNTA(ClaimFixtures[case_id])'],['Held-out cases','=COUNTIF(ClaimFixtures[split],"holdout")']]:dash.append(r)
chart=BarChart();chart.title='Source section changes';chart.add_data(Reference(dash,min_col=2,min_row=3,max_row=7));chart.set_categories(Reference(dash,min_col=1,min_row=3,max_row=7));chart.legend=None;chart.height=9;chart.width=18;dash.add_chart(chart,'D2')
for ws in wb:
 ws.freeze_panes='A2'
 # Excel rejects overlapping worksheet and table AutoFilters even though the
 # XML is well-formed. Structured tables already own their filter ranges.
 if not ws.tables:ws.auto_filter.ref=ws.dimensions
 for cell in ws[1]:cell.font=Font(name='Aptos',bold=True,color='FFFFFF',size=11);cell.fill=PatternFill('solid',fgColor='234D3B')
 for row in ws.iter_rows(min_row=2):
  for cell in row:cell.font=Font(name='Aptos',size=10);cell.alignment=Alignment(vertical='top')
 for col in ws.columns:
  letter=col[0].column_letter;ws.column_dimensions[letter].width=min(65,max(15,max(len(str(c.value or '')) for c in col[:60])+2))
 ws.sheet_view.showGridLines=False
readme.column_dimensions['A'].width=22;readme.column_dimensions['B'].width=110
for row in readme.iter_rows(min_row=2):
 row[1].alignment=Alignment(wrap_text=True,vertical='top')
 readme.row_dimensions[row[0].row].height=32
wb.calculation=CalcProperties(fullCalcOnLoad=True)
wb.save(out/'PharmaEvidence-Source-Tracker.xlsx')
print('Excel tracker and two CSV marts exported; formulas refer to the same source tables.')
