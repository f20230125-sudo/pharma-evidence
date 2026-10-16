import pathlib
import unittest
import openpyxl
import json

ROOT=pathlib.Path(__file__).resolve().parents[2]
class ExportChecks(unittest.TestCase):
 def test_excel_tables_do_not_overlap_worksheet_filters(self):
  workbook=openpyxl.load_workbook(ROOT/'public/downloads/PharmaEvidence-Source-Tracker.xlsx')
  for sheet in workbook:
   if sheet.tables:self.assertFalse(sheet.auto_filter.ref, sheet.title)
 def test_workbook_rows_reconcile_to_source_and_benchmark(self):
  workbook=openpyxl.load_workbook(ROOT/'public/downloads/PharmaEvidence-Source-Tracker.xlsx')
  corpus=json.loads((ROOT/'public/data/corpus.json').read_text(encoding='utf8'))
  self.assertEqual(workbook['Versions'].max_row-1,len(corpus['documents']))
  self.assertEqual(workbook['Section changes'].max_row-1,sum(len(c['sections']) for c in corpus['changes']))
  self.assertEqual(workbook['Benchmark fixtures'].max_row-1,120)
 def test_official_fda_reference_bytes_match_the_provenance(self):
  import hashlib
  for row in json.loads((ROOT/'data/fda-crosscheck.json').read_text()):
   data=(ROOT/'data/raw'/f"{row['product']}-fda-current.pdf").read_bytes()
   self.assertEqual(hashlib.sha256(data).hexdigest(),row['sha256'])
if __name__=='__main__':unittest.main()
